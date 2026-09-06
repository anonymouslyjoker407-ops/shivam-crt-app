from flask import Flask, render_template_string, request, redirect, url_for, session, Response
import json
import os
import base64
import requests
from datetime import datetime
import io

# Optional imports for Excel and PDF exports
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    EXCEL_SUPPORT = True
except ImportError:
    EXCEL_SUPPORT = False

try:
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    PDF_SUPPORT = True
except ImportError:
    PDF_SUPPORT = False

app = Flask(__name__)
app.secret_key = "shivam_crt_secure_master_key_2026"

GITHUB_TOKEN = os.environ.get("GITHUB_TOKEN")
GITHUB_REPO = os.environ.get("GITHUB_REPO")
GITHUB_BRANCH = os.environ.get("GITHUB_BRANCH", "main")
FILE_PATH = "shivam_crt_data.json"

def load_data():
    default_data = {
        "attendance": [],
        "expenses": [],
        "cash_expenses": []
    }
    
    if not GITHUB_TOKEN or not GITHUB_REPO:
        if not os.path.exists(FILE_PATH):
            with open(FILE_PATH, "w", encoding="utf-8") as f:
                json.dump(default_data, f, indent=4, ensure_ascii=False)
        try:
            with open(FILE_PATH, "r", encoding="utf-8") as f:
                d = json.load(f)
                for key in default_data:
                    if key not in d:
                        d[key] = default_data[key]
                return d
        except:
            return default_data

    url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{FILE_PATH}"
    headers = {"Authorization": f"Bearer {GITHUB_TOKEN}", "Accept": "vnd.github+json"}
    response = requests.get(url, headers=headers)
    
    if response.status_code == 200:
        try:
            file_content = response.json().get("content")
            decoded_content = base64.b64decode(file_content).decode("utf-8")
            d = json.loads(decoded_content)
            for key in default_data:
                if key not in d:
                    d[key] = default_data[key]
            return d
        except:
            return default_data
    else:
        save_data(default_data)
        return default_data

def save_data(data):
    if not GITHUB_TOKEN or not GITHUB_REPO:
        with open(FILE_PATH, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=4, ensure_ascii=False)
        return

    url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{FILE_PATH}"
    headers = {"Authorization": f"Bearer {GITHUB_TOKEN}", "Accept": "vnd.github+json"}
    
    get_res = requests.get(url, headers=headers)
    sha = get_res.json().get("sha") if get_res.status_code == 200 else None

    json_str = json.dumps(data, indent=4, ensure_ascii=False)
    encoded_content = base64.b64encode(json_str.encode("utf-8")).decode("utf-8")

    payload = {
        "message": "Auto-sync Shivam CRT master records",
        "content": encoded_content,
        "branch": GITHUB_BRANCH
    }
    if sha:
        payload["sha"] = sha

    requests.put(url, headers=headers, json=payload)

@app.route("/login", methods=["GET", "POST"])
def login():
    error = None
    if request.method == "POST":
        username = request.form.get("username")
        password = request.form.get("password")
        
        if username == "admin" and password == "admin123":
            session["user"] = username
            return redirect(url_for("index"))
        else:
            error = "Invalid credentials! Use admin / admin123"
            
    return render_template_string(LOGIN_HTML, error=error)

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))

@app.route("/", methods=["GET", "POST"])
def index():
    if "user" not in session:
        return redirect(url_for("login"))
    
    data = load_data()
    data.setdefault("attendance", [])
    data.setdefault("expenses", [])
    data.setdefault("cash_expenses", [])

    action = request.args.get("action", "dashboard")
    search_query = request.args.get("search", "").strip().lower()
    cash_search = request.args.get("cash_search", "").strip().lower()

    if request.method == "POST":
        form_type = request.form.get("form_type")
        
        if form_type == "add_attendance":
            attachments = []
            uploaded_files = request.files.getlist("attachments")
            for uploaded_file in uploaded_files:
                if uploaded_file and uploaded_file.filename:
                    file_name = uploaded_file.filename
                    file_bytes = uploaded_file.read()
                    file_data = base64.b64encode(file_bytes).decode("utf-8")
                    attachments.append({"file_name": file_name, "file_data": file_data})

            amt = request.form.get("amount", "0")
            conv = request.form.get("conveyance", "0")
            new_record = {
                "id": max([a["id"] for a in data["attendance"]], default=0) + 1,
                "date": request.form.get("date", datetime.now().strftime("%Y-%m-%d")),
                "name": request.form.get("name", ""),
                "card_number": request.form.get("card_number", ""),
                "role": request.form.get("role", ""),
                "in_time": request.form.get("in_time", ""),
                "out_time": request.form.get("out_time", ""),
                "signature": request.form.get("signature", "Yes"),
                "description": request.form.get("description", ""),
                "amount": int(amt) if amt.isdigit() else 0,
                "conveyance": int(conv) if conv.isdigit() else 0,
                "attachments": attachments
            }
            data["attendance"].append(new_record)
            save_data(data)
            return redirect(url_for("index", action="attendance"))

        elif form_type == "edit_attendance":
            a_id = int(request.form.get("record_id", 0))
            for a in data["attendance"]:
                if a["id"] == a_id:
                    a["date"] = request.form.get("date", a.get("date", ""))
                    a["name"] = request.form.get("name", "")
                    a["card_number"] = request.form.get("card_number", "")
                    a["role"] = request.form.get("role", "")
                    a["in_time"] = request.form.get("in_time", "")
                    a["out_time"] = request.form.get("out_time", "")
                    a["signature"] = request.form.get("signature", "Yes")
                    a["description"] = request.form.get("description", "")
                    amt = request.form.get("amount", "0")
                    conv = request.form.get("conveyance", "0")
                    a["amount"] = int(amt) if amt.isdigit() else 0
                    a["conveyance"] = int(conv) if conv.isdigit() else 0
                    
                    existing_atts = a.get("attachments", [])
                    retained_indices = request.form.getlist("keep_attachments")
                    retained_atts = [existing_atts[int(i)] for i in retained_indices if int(i) < len(existing_atts)]

                    uploaded_files = request.files.getlist("attachments")
                    new_attachments = []
                    for uploaded_file in uploaded_files:
                        if uploaded_file and uploaded_file.filename:
                            file_name = uploaded_file.filename
                            file_bytes = uploaded_file.read()
                            file_data = base64.b64encode(file_bytes).decode("utf-8")
                            new_attachments.append({"file_name": file_name, "file_data": file_data})
                    
                    a["attachments"] = retained_atts + new_attachments
            save_data(data)
            return redirect(url_for("index", action="attendance"))

        elif form_type == "add_expense":
            amt = request.form.get("amount", "0")
            new_exp = {
                "id": max([e["id"] for e in data["expenses"]], default=0) + 1,
                "category": request.form.get("category", ""),
                "amount": int(amt) if amt.isdigit() else 0,
                "date": request.form.get("date", "")
            }
            data["expenses"].append(new_exp)
            save_data(data)
            return redirect(url_for("index", action="expenses"))

        elif form_type == "add_cash_expense":
            attachments = []
            uploaded_files = request.files.getlist("attachments")
            for uploaded_file in uploaded_files:
                if uploaded_file and uploaded_file.filename:
                    file_name = uploaded_file.filename
                    file_bytes = uploaded_file.read()
                    file_data = base64.b64encode(file_bytes).decode("utf-8")
                    attachments.append({"file_name": file_name, "file_data": file_data})

            amt = request.form.get("amount", "0")
            new_cash_exp = {
                "id": max([c["id"] for c in data["cash_expenses"]], default=0) + 1,
                "date": request.form.get("date", ""),
                "pay_to": request.form.get("pay_to", ""),
                "description": request.form.get("description", ""),
                "amount": int(amt) if amt.isdigit() else 0,
                "attachments": attachments
            }
            data["cash_expenses"].append(new_cash_exp)
            save_data(data)
            return redirect(url_for("index", action="cash_expenses"))

        elif form_type == "edit_cash_expense":
            c_id = int(request.form.get("record_id", 0))
            for c in data["cash_expenses"]:
                if c["id"] == c_id:
                    c["date"] = request.form.get("date", "")
                    c["pay_to"] = request.form.get("pay_to", "")
                    c["description"] = request.form.get("description", "")
                    amt = request.form.get("amount", "0")
                    c["amount"] = int(amt) if amt.isdigit() else 0
                    
                    existing_atts = c.get("attachments", [])
                    retained_indices = request.form.getlist("keep_attachments")
                    retained_atts = [existing_atts[int(i)] for i in retained_indices if int(i) < len(existing_atts)]

                    uploaded_files = request.files.getlist("attachments")
                    new_attachments = []
                    for uploaded_file in uploaded_files:
                        if uploaded_file and uploaded_file.filename:
                            file_name = uploaded_file.filename
                            file_bytes = uploaded_file.read()
                            file_data = base64.b64encode(file_bytes).decode("utf-8")
                            new_attachments.append({"file_name": file_name, "file_data": file_data})
                    
                    c["attachments"] = retained_atts + new_attachments
            save_data(data)
            return redirect(url_for("index", action="cash_expenses"))

    # Filtering attendance
    filtered_attendance = data["attendance"]
    if search_query:
        filtered_attendance = [
            a for a in data["attendance"] 
            if search_query in str(a.get("name", "")).lower() or
               search_query in str(a.get("card_number", "")).lower() or
               search_query in str(a.get("role", "")).lower() or
               search_query in str(a.get("date", "")).lower() or
               search_query in str(a.get("description", "")).lower() or
               search_query in str(a.get("amount", "")).lower() or
               search_query in str(a.get("conveyance", "")).lower()
        ]

    # Filtering cash expenses
    filtered_cash_expenses = data["cash_expenses"]
    if cash_search:
        filtered_cash_expenses = [
            c for c in data["cash_expenses"]
            if cash_search in str(c.get("date", "")).lower() or
               cash_search in str(c.get("pay_to", "")).lower() or
               cash_search in str(c.get("description", "")).lower() or
               cash_search in str(c.get("amount", "")).lower()
        ]

    total_amount = sum(a.get("amount", 0) for a in data["attendance"])
    total_conveyance = sum(a.get("conveyance", 0) for a in data["attendance"])
    total_expenses = sum(e.get("amount", 0) for e in data["expenses"])
    total_cash_expenses = sum(c.get("amount", 0) for c in data["cash_expenses"])

    filtered_total_amount = sum(a.get("amount", 0) for a in filtered_attendance)
    filtered_total_conveyance = sum(a.get("conveyance", 0) for a in filtered_attendance)
    filtered_grand_total = filtered_total_amount + filtered_total_conveyance
    filtered_cash_total = sum(c.get("amount", 0) for c in filtered_cash_expenses)

    return render_template_string(
        DASHBOARD_HTML,
        data=data,
        filtered_attendance=filtered_attendance,
        filtered_cash_expenses=filtered_cash_expenses,
        action=action,
        search_query=search_query,
        cash_search=cash_search,
        sources_status="GitHub API Synced" if (GITHUB_TOKEN and GITHUB_REPO) else "Local Storage Mode",
        total_amount=total_amount,
        total_conveyance=total_conveyance,
        total_expenses=total_expenses,
        total_cash_expenses=total_cash_expenses,
        filtered_total_amount=filtered_total_amount,
        filtered_total_conveyance=filtered_total_conveyance,
        filtered_grand_total=filtered_grand_total,
        filtered_cash_total=filtered_cash_total
    )

@app.route("/export/excel")
def export_excel():
    if "user" not in session:
        return redirect(url_for("login"))
    data = load_data()
    
    if not EXCEL_SUPPORT:
        return "openpyxl library not installed on server.", 400

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Ledger"

    headers = ["Sr #", "Date", "Name", "Card No", "Role", "In Time", "Out Time", "Signature", "Description", "Amount (₹)", "Conveyance (₹)", "Total (₹)"]
    ws.append(headers)

    for idx, a in enumerate(data.get("attendance", []), 1):
        ws.append([
            idx,
            a.get("date", ""),
            a.get("name", ""),
            a.get("card_number", ""),
            a.get("role", ""),
            a.get("in_time", ""),
            a.get("out_time", ""),
            a.get("signature", ""),
            a.get("description", ""),
            a.get("amount", 0),
            a.get("conveyance", 0),
            a.get("amount", 0) + a.get("conveyance", 0)
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return Response(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment;filename=Shivam_CRT_Attendance_Ledger.xlsx"}
    )

@app.route("/export/pdf")
def export_pdf():
    if "user" not in session:
        return redirect(url_for("login"))
    data = load_data()

    if not PDF_SUPPORT:
        return "reportlab library not installed on server.", 400

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20)
    elements = []

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=16, textColor=colors.HexColor('#6366f1'), spaceAfter=12)

    elements.append(Paragraph("Shivam CRT - Attendance & Operations Master Ledger", title_style))
    elements.append(Spacer(1, 10))

    table_data = [["Sr", "Date", "Name", "Card", "Role", "In", "Out", "Sig", "Desc", "Amount", "Conv", "Total"]]
    for idx, a in enumerate(data.get("attendance", []), 1):
        table_data.append([
            str(idx),
            str(a.get("date", "")),
            str(a.get("name", "")),
            str(a.get("card_number", "")),
            str(a.get("role", "")),
            str(a.get("in_time", "")),
            str(a.get("out_time", "")),
            str(a.get("signature", "")),
            str(a.get("description", "")),
            f"Rs {a.get('amount', 0)}",
            f"Rs {a.get('conveyance', 0)}",
            f"Rs {a.get('amount', 0) + a.get('conveyance', 0)}"
        ])

    t = Table(table_data, colWidths=[25, 65, 95, 60, 80, 50, 50, 40, 90, 60, 60, 65])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e1b4b')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,0), 8),
        ('BOTTOMPADDING', (0,0), (-1,0), 5),
        ('BACKGROUND', (0,1), (-1,-1), colors.HexColor('#f8fafc')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('FONTNAME', (0,1), (-1,-1), 'Helvetica'),
        ('FONTSIZE', (0,1), (-1,-1), 7),
    ]))

    elements.append(t)
    doc.build(elements)
    buffer.seek(0)

    return Response(
        buffer,
        mimetype="application/pdf",
        headers={"Content-Disposition": "attachment;filename=Shivam_CRT_Attendance_Ledger.pdf"}
    )

@app.route("/export/cash_excel")
def export_cash_excel():
    if "user" not in session:
        return redirect(url_for("login"))
    data = load_data()
    
    if not EXCEL_SUPPORT:
        return "openpyxl library not installed on server.", 400

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cash Expenses"

    headers = ["Sr #", "Date", "Pay To", "Description", "Amount (₹)", "Attachments Info"]
    ws.append(headers)

    for idx, c in enumerate(data.get("cash_expenses", []), 1):
        att_names = ", ".join([att.get("file_name", "") for att in c.get("attachments", [])])
        ws.append([
            idx,
            c.get("date", ""),
            c.get("pay_to", ""),
            c.get("description", ""),
            c.get("amount", 0),
            att_names if att_names else "No Files"
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return Response(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment;filename=Shivam_CRT_Cash_Expenses.xlsx"}
    )

@app.route("/export/cash_pdf")
def export_cash_pdf():
    if "user" not in session:
        return redirect(url_for("login"))
    data = load_data()

    if not PDF_SUPPORT:
        return "reportlab library not installed on server.", 400

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20)
    elements = []

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=16, textColor=colors.HexColor('#10b981'), spaceAfter=12)

    elements.append(Paragraph("Shivam CRT - Cash Expenses Master Ledger", title_style))
    elements.append(Spacer(1, 10))

    # Group cash expenses by date
    cash_expenses = data.get("cash_expenses", [])
    
    # Sort or iterate grouped by date
    # Let's organize data by date groups
    grouped_cash = {}
    for c in cash_expenses:
        d = c.get("date", "Unspecified Date")
        if d not in grouped_cash:
            grouped_cash[d] = []
        grouped_cash[d].append(c)

    for d, items in grouped_cash.items():
        elements.append(Paragraph(f"<b>Date Group: {d}</b>", ParagraphStyle('DateGroupStyle', parent=styles['Heading3'], fontSize=11, textColor=colors.HexColor('#064e3b'), spaceBefore=8, spaceAfter=4)))
        
        table_data = [["Sr", "Pay To", "Description", "Amount", "Attachments & Images"]]
        date_total = 0
        
        for idx, c in enumerate(items, 1):
            date_total += c.get('amount', 0)
            
            # Format attachments column content (text + images if possible)
            att_flowables = []
            atts = c.get("attachments", [])
            if atts:
                for att in atts:
                    fname = att.get("file_name", "")
                    att_flowables.append(Paragraph(f"• {fname}", ParagraphStyle('AttText', fontSize=7, textColor=colors.HexColor('#1e293b'))))
                    # Try embedding image if it's an image
                    fdata = att.get("file_data", "")
                    if fdata and (fname.lower().endswith(('.png', '.jpg', '.jpeg', '.gif', '.webp'))):
                        try:
                            img_bytes = base64.b64decode(fdata)
                            img_io = io.BytesIO(img_bytes)
                            rl_img = RLImage(img_io, width=80, height=60)
                            att_flowables.append(rl_img)
                            att_flowables.append(Spacer(1, 4))
                        except:
                            pass
            else:
                att_flowables = [Paragraph("No Files", ParagraphStyle('NoAtt', fontSize=7, textColor=colors.HexColor('#64748b')))]

            table_data.append([
                str(idx),
                str(c.get("pay_to", "")),
                str(c.get("description", "")),
                f"Rs {c.get('amount', 0)}",
                att_flowables
            ])

        t = Table(table_data, colWidths=[25, 110, 150, 75, 160])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#064e3b')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 8),
            ('BOTTOMPADDING', (0,0), (-1,0), 5),
            ('BACKGROUND', (0,1), (-1,-1), colors.HexColor('#f0fdf4')),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#a7f3d0')),
            ('FONTNAME', (0,1), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,1), (-1,-1), 7),
        ]))
        elements.append(t)
        
        # Subtotal paragraph for the date
        elements.append(Paragraph(f"<b>Subtotal for {d}: Rs {date_total}</b>", ParagraphStyle('SubTotalStyle', parent=styles['Normal'], fontSize=9, textColor=colors.HexColor('#047857'), alignment=2, spaceBefore=4, spaceAfter=10)))
        elements.append(Spacer(1, 5))

    doc.build(elements)
    buffer.seek(0)

    return Response(
        buffer,
        mimetype="application/pdf",
        headers={"Content-Disposition": "attachment;filename=Shivam_CRT_Cash_Expenses.pdf"}
    )

@app.route("/delete/<string:category>/<int:item_id>")
def delete_item(category, item_id):
    if "user" not in session:
        return redirect(url_for("login"))

    data = load_data()
    data.setdefault("attendance", [])
    data.setdefault("expenses", [])
    data.setdefault("cash_expenses", [])

    if category == "attendance":
        data["attendance"] = [a for a in data["attendance"] if a["id"] != item_id]
        redirect_action = "attendance"
    elif category == "expense":
        data["expenses"] = [e for e in data["expenses"] if e["id"] != item_id]
        redirect_action = "expenses"
    elif category == "cash_expense":
        data["cash_expenses"] = [c for c in data["cash_expenses"] if c["id"] != item_id]
        redirect_action = "cash_expenses"
    
    save_data(data)
    return redirect(url_for("index", action=redirect_action))

LOGIN_HTML = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Shivam CRT - Login</title>
    <script src="https://cdn.tailwindcss.com"></script>
</head>
<body class="bg-gray-950 text-white flex items-center justify-center h-screen px-4">
    <div class="bg-gray-900 p-8 rounded-2xl shadow-2xl w-full max-w-md border border-gray-800">
        <h2 class="text-3xl font-black text-center mb-2 text-indigo-400">⚡ Shivam CRT</h2>
        <p class="text-xs text-gray-400 text-center mb-6">Master Attendance & Payroll Management Suite</p>
        
        {% if error %}
            <div class="bg-red-500/20 border border-red-500 text-red-300 p-3 rounded-xl mb-4 text-xs text-center">{{ error }}</div>
        {% endif %}
        
        <form method="POST" class="space-y-4">
            <div>
                <label class="block text-xs text-gray-400 mb-1">Username</label>
                <input type="text" name="username" placeholder="admin" class="w-full px-4 py-2.5 bg-gray-800 rounded-xl border border-gray-700 focus:outline-none focus:ring-2 focus:ring-indigo-400 text-sm">
            </div>
            <div>
                <label class="block text-xs text-gray-400 mb-1">Password</label>
                <input type="password" name="password" placeholder="password" class="w-full px-4 py-2.5 bg-gray-800 rounded-xl border border-gray-700 focus:outline-none focus:ring-2 focus:ring-indigo-400 text-sm">
            </div>
            <button type="submit" class="w-full py-3 bg-indigo-500 hover:bg-indigo-600 transition rounded-xl font-bold text-gray-950 shadow-lg text-sm mt-2">Log In</button>
        </form>

        <div class="mt-6 p-4 bg-gray-800/50 rounded-xl border border-gray-800 text-xs space-y-1 text-gray-400 text-center">
            <p><strong class="text-indigo-400">👑 Credentials:</strong> admin / admin123</p>
        </div>
    </div>
</body>
</html>
"""

DASHBOARD_HTML = """
<!DOCTYPE html>
<html lang="en" id="htmlRoot">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Shivam CRT Suite</title>
    <script src="https://cdn.tailwindcss.com"></script>
</head>
<body class="bg-gray-950 text-gray-100 font-sans transition-colors duration-200" id="bodyTheme">
    <div class="flex h-screen overflow-hidden">
        
        <!-- SIDEBAR -->
        <div class="hidden md:flex flex-col w-64 bg-gray-900 border-r border-gray-800 p-6 sidebar-panel" id="sidebarPanel">
            <h1 class="text-2xl font-black text-indigo-400 mb-1 tracking-wider">⚡ Shivam CRT</h1>
            <p class="text-xs text-gray-400 mb-6 font-mono">Operations Portal</p>
            <div class="mb-6 px-3 py-1.5 bg-indigo-500/10 border border-indigo-500/30 rounded-lg text-[10px] text-indigo-400 font-mono text-center">
                🟢 {{ sources_status }}
            </div>
            
            <nav class="space-y-2 flex-1">
                <a href="/?action=dashboard" class="block py-2.5 px-4 rounded-xl font-semibold transition {% if action == 'dashboard' %}bg-indigo-500/10 text-indigo-400{% else %}text-gray-400 hover:bg-gray-800{% endif %}">📊 Dashboard</a>
                <a href="/?action=attendance" class="block py-2.5 px-4 rounded-xl font-semibold transition {% if action == 'attendance' %}bg-indigo-500/10 text-indigo-400{% else %}text-gray-400 hover:bg-gray-800{% endif %}">📋 Attendance & Ledger</a>
                <a href="/?action=expenses" class="block py-2.5 px-4 rounded-xl font-semibold transition {% if action == 'expenses' %}bg-indigo-500/10 text-indigo-400{% else %}text-gray-400 hover:bg-gray-800{% endif %}">💡 Expenses Ledger</a>
                <a href="/?action=cash_expenses" class="block py-2.5 px-4 rounded-xl font-semibold transition {% if action == 'cash_expenses' %}bg-indigo-500/10 text-indigo-400{% else %}text-gray-400 hover:bg-gray-800{% endif %}">💵 Cash Expense</a>
                <a href="/logout" class="block py-2.5 px-4 rounded-xl font-semibold text-red-400 hover:bg-red-500/10 transition mt-8">🚪 Log Out</a>
            </nav>
        </div>

        <!-- MAIN CONTAINER -->
        <div class="flex-1 flex flex-col overflow-y-auto">
            <header class="bg-gray-900 border-b border-gray-800 p-4 flex justify-between items-center header-panel" id="headerPanel">
                <h1 class="text-lg font-black text-indigo-400">⚡ Shivam CRT</h1>
                <div class="flex items-center space-x-3">
                    <button onclick="toggleTheme()" class="px-3 py-1.5 bg-gray-800 hover:bg-gray-700 text-indigo-400 rounded-xl text-xs font-bold border border-gray-700 transition" id="themeToggleBtn">☀️ Light Mode</button>
                    <a href="/logout" class="text-xs text-red-400 font-bold bg-red-500/10 px-3 py-1.5 rounded-lg md:hidden">Log Out</a>
                </div>
            </header>

            <!-- MOBILE NAV -->
            <div class="flex md:hidden bg-gray-900 p-2 overflow-x-auto space-x-2 border-b border-gray-800 shrink-0">
                <a href="/?action=dashboard" class="px-3 py-1.5 text-xs font-semibold rounded-lg {% if action == 'dashboard' %}bg-indigo-500 text-gray-950{% else %}bg-gray-800 text-gray-300{% endif %} whitespace-nowrap">Dashboard</a>
                <a href="/?action=attendance" class="px-3 py-1.5 text-xs font-semibold rounded-lg {% if action == 'attendance' %}bg-indigo-500 text-gray-950{% else %}bg-gray-800 text-gray-300{% endif %} whitespace-nowrap">Attendance</a>
                <a href="/?action=expenses" class="px-3 py-1.5 text-xs font-semibold rounded-lg {% if action == 'expenses' %}bg-indigo-500 text-gray-950{% else %}bg-gray-800 text-gray-300{% endif %} whitespace-nowrap">Expenses</a>
                <a href="/?action=cash_expenses" class="px-3 py-1.5 text-xs font-semibold rounded-lg {% if action == 'cash_expenses' %}bg-indigo-500 text-gray-950{% else %}bg-gray-800 text-gray-300{% endif %} whitespace-nowrap">Cash Expense</a>
            </div>

            <div class="p-4 sm:p-6 max-w-7xl mx-auto w-full space-y-6">
                
                {% if action == 'dashboard' %}
                <div class="space-y-6">
                    <h2 class="text-2xl font-black dynamic-text">📊 Master Operational Dashboard</h2>
                    
                    <div class="grid grid-cols-1 sm:grid-cols-4 gap-4">
                        <div class="bg-gray-900 p-5 rounded-2xl border border-gray-800 shadow-xl card-panel">
                            <p class="text-xs text-gray-400 font-bold">Total Amount Allocated</p>
                            <h3 class="text-xl font-black text-indigo-400 mt-1">₹{{ total_amount }}</h3>
                        </div>
                        <div class="bg-gray-900 p-5 rounded-2xl border border-gray-800 shadow-xl card-panel">
                            <p class="text-xs text-gray-400 font-bold">Total Conveyance</p>
                            <h3 class="text-xl font-black text-amber-400 mt-1">₹{{ total_conveyance }}</h3>
                        </div>
                        <div class="bg-gray-900 p-5 rounded-2xl border border-gray-800 shadow-xl card-panel">
                            <p class="text-xs text-gray-400 font-bold">Total Expenses</p>
                            <h3 class="text-xl font-black text-red-400 mt-1">₹{{ total_expenses }}</h3>
                        </div>
                        <div class="bg-gray-900 p-5 rounded-2xl border border-gray-800 shadow-xl card-panel">
                            <p class="text-xs text-gray-400 font-bold">Total Cash Expenses</p>
                            <h3 class="text-xl font-black text-emerald-400 mt-1">₹{{ total_cash_expenses }}</h3>
                        </div>
                    </div>

                    <div class="bg-gray-900 p-6 rounded-2xl border border-gray-800 shadow-xl card-panel">
                        <h3 class="text-lg font-bold dynamic-text mb-4">🚀 Quick System Overview</h3>
                        <ul class="space-y-3 text-sm text-gray-300">
                            <li class="flex justify-between p-3 bg-gray-800/40 rounded-xl"><span>Total Attendance Records:</span> <strong class="text-indigo-400">{{ data.attendance|length }}</strong></li>
                            <li class="flex justify-between p-3 bg-gray-800/40 rounded-xl"><span>Total Expense Categories:</span> <strong class="text-red-400">{{ data.expenses|length }}</strong></li>
                            <li class="flex justify-between p-3 bg-gray-800/40 rounded-xl"><span>Total Cash Expense Records:</span> <strong class="text-emerald-400">{{ data.cash_expenses|length }}</strong></li>
                        </ul>
                    </div>
                </div>

                {% elif action == 'attendance' %}
                <div class="space-y-6">
                    <div class="bg-gray-900 p-6 rounded-2xl border border-gray-800 shadow-xl card-panel">
                        <h2 class="text-xl font-bold text-indigo-400 mb-4">➕ Add Attendance & Operations Record</h2>
                        <form method="POST" enctype="multipart/form-data" class="grid grid-cols-1 sm:grid-cols-3 gap-4">
                            <input type="hidden" name="form_type" value="add_attendance">
                            <div>
                                <label class="text-xs text-gray-400">Date</label>
                                <input type="date" name="date" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input">
                            </div>
                            <div>
                                <label class="text-xs text-gray-400">Full Name</label>
                                <input type="text" name="name" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input">
                            </div>
                            <div>
                                <label class="text-xs text-gray-400">Card Number</label>
                                <input type="text" name="card_number" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input" placeholder="CRT-XXXX">
                            </div>
                            <div>
                                <label class="text-xs text-gray-400">Role / Designation</label>
                                <input type="text" name="role" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input" placeholder="Trainer / Analyst">
                            </div>
                            <div>
                                <label class="text-xs text-gray-400">In Time</label>
                                <input type="text" name="in_time" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input" placeholder="09:00 AM">
                            </div>
                            <div>
                                <label class="text-xs text-gray-400">Out Time</label>
                                <input type="text" name="out_time" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input" placeholder="06:00 PM">
                            </div>
                            <div>
                                <label class="text-xs text-gray-400">Signature Status</label>
                                <select name="signature" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input">
                                    <option value="Yes">Yes</option>
                                    <option value="No">No</option>
                                    <option value="Mistake">Mistake</option>
                                </select>
                            </div>
                            <div>
                                <label class="text-xs text-gray-400">Amount (₹)</label>
                                <input type="number" name="amount" value="0" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input">
                            </div>
                            <div>
                                <label class="text-xs text-gray-400">Conveyance (₹)</label>
                                <input type="number" name="conveyance" value="0" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input">
                            </div>
                            <div class="sm:col-span-2">
                                <label class="text-xs text-gray-400">Description / Note</label>
                                <input type="text" name="description" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input" placeholder="Task details or notes...">
                            </div>
                            <div>
                                <label class="text-xs text-gray-400">Upload Multiple Files (PDF/Images)</label>
                                <input type="file" name="attachments" multiple accept=".pdf,image/*" class="w-full mt-1 p-2 bg-gray-800 rounded-xl border border-gray-700 text-xs text-gray-300">
                            </div>
                            <div class="sm:col-span-3">
                                <button type="submit" class="w-full py-3 bg-indigo-500 hover:bg-indigo-600 font-bold text-gray-950 rounded-xl transition text-sm shadow-lg">Save Attendance Record</button>
                            </div>
                        </form>
                    </div>

                    <!-- ATTENDANCE TABLE WITH DATE-WISE GROUPING -->
                    <div class="bg-gray-900 rounded-2xl border border-gray-800 shadow-xl overflow-hidden card-panel">
                        <div class="p-4 border-b border-gray-800 flex flex-col sm:flex-row justify-between items-center gap-4">
                            <h3 class="font-bold dynamic-text">📋 Date-Wise Grouped Attendance Ledger</h3>
                            
                            <div class="flex flex-wrap items-center gap-2 w-full sm:w-auto justify-end">
                                <form method="GET" action="/" class="flex items-center gap-2 w-full sm:w-auto">
                                    <input type="hidden" name="action" value="attendance">
                                    <input type="text" name="search" value="{{ search_query }}" placeholder="Search record..." class="px-3 py-1.5 bg-gray-800 border border-gray-700 rounded-xl text-xs text-white focus:outline-none focus:ring-2 focus:ring-indigo-400 w-full sm:w-64">
                                    <button type="submit" class="px-3 py-1.5 bg-indigo-500 text-gray-950 font-bold rounded-xl text-xs hover:bg-indigo-600 transition">Search</button>
                                    {% if search_query %}
                                        <a href="/?action=attendance" class="px-2 py-1.5 bg-gray-700 text-gray-300 rounded-xl text-xs hover:bg-gray-600">Clear</a>
                                    {% endif %}
                                </form>
                                <div class="flex items-center gap-1 border-l border-gray-700 pl-2">
                                    <a href="/export/excel" class="px-3 py-1.5 bg-emerald-600 hover:bg-emerald-700 text-white font-bold rounded-xl text-xs transition flex items-center gap-1">📊 Excel</a>
                                    <a href="/export/pdf" class="px-3 py-1.5 bg-red-600 hover:bg-red-700 text-white font-bold rounded-xl text-xs transition flex items-center gap-1">📄 PDF</a>
                                </div>
                            </div>
                        </div>

                        <div class="overflow-x-auto">
                            <table class="w-full text-left border-collapse text-xs">
                                <thead>
                                    <tr class="bg-gray-800/80 text-gray-300 uppercase tracking-wider font-mono">
                                        <th class="p-3 border-r border-gray-700">Sr #</th>
                                        <th class="p-3 border-r border-gray-700">Date</th>
                                        <th class="p-3 border-r border-gray-700">Name</th>
                                        <th class="p-3 border-r border-gray-700">Card No</th>
                                        <th class="p-3 border-r border-gray-700">Role</th>
                                        <th class="p-3 border-r border-gray-700">In / Out</th>
                                        <th class="p-3 border-r border-gray-700">Sig</th>
                                        <th class="p-3 border-r border-gray-700">Description</th>
                                        <th class="p-3 border-r border-gray-700">Amount</th>
                                        <th class="p-3 border-r border-gray-700">Conveyance</th>
                                        <th class="p-3 border-r border-gray-700">Total</th>
                                        <th class="p-3 border-r border-gray-700">Files</th>
                                        <th class="p-3 text-center">Actions</th>
                                    </tr>
                                </thead>
                                <tbody class="divide-y divide-gray-800 font-mono">
                                    {% set ns = namespace(current_date='', date_amt=0, date_conv=0) %}
                                    {% for a in filtered_attendance %}
                                        {% if a.date != ns.current_date %}
                                            {% if not loop.first %}
                                            <!-- Date Subtotal Row -->
                                            <tr class="bg-indigo-950/40 font-bold border-t-2 border-indigo-500/40">
                                                <td colspan="8" class="p-2.5 text-right text-indigo-300">Subtotal for {{ ns.current_date }}:</td>
                                                <td class="p-2.5 text-gray-200">₹{{ ns.date_amt }}</td>
                                                <td class="p-2.5 text-amber-400">₹{{ ns.date_conv }}</td>
                                                <td colspan="3" class="p-2.5 text-emerald-400">₹{{ ns.date_amt + ns.date_conv }}</td>
                                            </tr>
                                            {% endif %}
                                            {% set ns.current_date = a.date %}
                                            {% set ns.date_amt = a.amount %}
                                            {% set ns.date_conv = a.conveyance %}
                                            <!-- Date Group Header -->
                                            <tr class="bg-indigo-900/40 text-indigo-400 font-bold">
                                                <td colspan="13" class="p-2.5 px-4 text-xs tracking-wider">📅 Date Group: {{ a.date or 'Unspecified Date' }}</td>
                                            </tr>
                                        {% else %}
                                            {% set ns.date_amt = ns.date_amt + a.amount %}
                                            {% set ns.date_conv = ns.date_conv + a.conveyance %}
                                        {% endif %}

                                    <tr class="hover:bg-gray-800/30 transition">
                                        <td class="p-3 border-r border-gray-800">{{ loop.index }}</td>
                                        <td class="p-3 border-r border-gray-800 text-gray-300">{{ a.date }}</td>
                                        <td class="p-3 border-r border-gray-800 font-bold text-indigo-400">{{ a.name }}</td>
                                        <td class="p-3 border-r border-gray-800">{{ a.card_number }}</td>
                                        <td class="p-3 border-r border-gray-800">{{ a.role }}</td>
                                        <td class="p-3 border-r border-gray-800 text-[10px]">{{ a.in_time }} - {{ a.out_time }}</td>
                                        <td class="p-3 border-r border-gray-800">
                                            <span class="px-2 py-0.5 rounded text-[10px] {% if a.signature == 'Yes' %}bg-emerald-500/10 text-emerald-400{% elif a.signature == 'No' %}bg-red-500/10 text-red-400{% else %}bg-amber-500/10 text-amber-400{% endif %}">{{ a.signature }}</span>
                                        </td>
                                        <td class="p-3 border-r border-gray-800 text-gray-300">{{ a.description }}</td>
                                        <td class="p-3 border-r border-gray-800 text-gray-200">₹{{ a.amount }}</td>
                                        <td class="p-3 border-r border-gray-800 text-amber-400">₹{{ a.conveyance }}</td>
                                        <td class="p-3 border-r border-gray-800 font-bold text-emerald-400">₹{{ a.amount + a.conveyance }}</td>
                                        <td class="p-3 border-r border-gray-800">
                                            {% if a.attachments %}
                                                <div class="flex flex-col space-y-1">
                                                    {% for att in a.attachments %}
                                                    <button onclick="openFileViewer('data:application/octet-stream;base64,{{ att.file_data }}', '{{ att.file_name }}')" class="text-indigo-400 underline hover:text-indigo-300 text-[10px] text-left">📎 {{ att.file_name[:12] }}...</button>
                                                    {% endfor %}
                                                </div>
                                            {% else %}
                                                <span class="text-gray-500 text-[10px]">No Files</span>
                                            {% endif %}
                                        </td>
                                        <td class="p-3 text-center space-x-2 whitespace-nowrap">
                                            <button onclick='openEditModal({{ a|tojson|safe }})' class="text-blue-400 bg-blue-500/10 px-2 py-1 rounded text-[10px] font-bold cursor-pointer">Edit</button>
                                            <a href="/delete/attendance/{{ a.id }}" onclick="return confirm('Confirm delete record?');" class="text-red-400 bg-red-500/10 px-2 py-1 rounded text-[10px] font-bold inline-block">Delete</a>
                                        </td>
                                    </tr>
                                        {% if loop.last and filtered_attendance %}
                                        <!-- Last Date Subtotal Row -->
                                        <tr class="bg-indigo-950/40 font-bold border-t-2 border-indigo-500/40">
                                            <td colspan="8" class="p-2.5 text-right text-indigo-300">Subtotal for {{ ns.current_date }}:</td>
                                            <td class="p-2.5 text-gray-200">₹{{ ns.date_amt }}</td>
                                            <td class="p-2.5 text-amber-400">₹{{ ns.date_conv }}</td>
                                            <td colspan="3" class="p-2.5 text-emerald-400">₹{{ ns.date_amt + ns.date_conv }}</td>
                                        </tr>
                                        {% endif %}
                                    {% endfor %}
                                </tbody>
                                <tfoot>
                                    <tr class="bg-gray-800/90 font-mono font-bold text-gray-200 border-t-2 border-gray-700">
                                        <td colspan="8" class="p-3 text-right uppercase tracking-wider text-indigo-400">Grand Total Sum:</td>
                                        <td class="p-3 border-r border-gray-700 text-gray-200">₹{{ filtered_total_amount }}</td>
                                        <td class="p-3 border-r border-gray-700 text-amber-400">₹{{ filtered_total_conveyance }}</td>
                                        <td colspan="3" class="p-3 text-emerald-400">₹{{ filtered_grand_total }}</td>
                                    </tr>
                                </tfoot>
                            </table>
                        </div>
                    </div>
                </div>

                {% elif action == 'expenses' %}
                <div class="space-y-6">
                    <div class="bg-gray-900 p-6 rounded-2xl border border-gray-800 shadow-xl card-panel">
                        <h2 class="text-xl font-bold text-red-400 mb-4">💡 Expenses Ledger</h2>
                        <form method="POST" class="grid grid-cols-1 sm:grid-cols-3 gap-4">
                            <input type="hidden" name="form_type" value="add_expense">
                            <div>
                                <label class="text-xs text-gray-400">Category</label>
                                <input type="text" name="category" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input">
                            </div>
                            <div>
                                <label class="text-xs text-gray-400">Amount (₹)</label>
                                <input type="number" name="amount" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input">
                            </div>
                            <div>
                                <label class="text-xs text-gray-400">Date</label>
                                <input type="date" name="date" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input">
                            </div>
                            <div class="sm:col-span-3">
                                <button type="submit" class="w-full py-3 bg-red-500 hover:bg-red-600 font-bold text-gray-950 rounded-xl transition text-sm shadow-lg">Save Expense</button>
                            </div>
                        </form>
                    </div>

                    <div class="bg-gray-900 rounded-2xl border border-gray-800 shadow-xl overflow-hidden card-panel">
                        <div class="p-4 border-b border-gray-800"><h3 class="font-bold dynamic-text">📉 Expenses Ledger</h3></div>
                        <div class="overflow-x-auto">
                            <table class="w-full text-left border-collapse text-sm">
                                <thead>
                                    <tr class="bg-gray-800/50 text-gray-400 font-mono text-xs">
                                        <th class="p-3">Description</th>
                                        <th class="p-3">Date</th>
                                        <th class="p-3">Amount</th>
                                        <th class="p-3 text-center">Actions</th>
                                    </tr>
                                </thead>
                                <tbody class="divide-y divide-gray-800 font-mono text-xs">
                                    {% for e in data.expenses %}
                                    <tr class="hover:bg-gray-800/30">
                                        <td class="p-3 font-semibold text-indigo-400">{{ e.category }}</td>
                                        <td class="p-3 text-gray-300">{{ e.date }}</td>
                                        <td class="p-3 text-red-400 font-bold">₹{{ e.amount }}</td>
                                        <td class="p-3 text-center">
                                            <a href="/delete/expense/{{ e.id }}" onclick="return confirm('Confirm delete expense?');" class="text-red-400 bg-red-500/10 px-2 py-1 rounded text-[10px] font-bold">Delete</a>
                                        </td>
                                    </tr>
                                    {% endfor %}
                                </tbody>
                            </table>
                        </div>
                    </div>
                </div>

                {% elif action == 'cash_expenses' %}
                <div class="space-y-6">
                    <div class="bg-gray-900 p-6 rounded-2xl border border-gray-800 shadow-xl card-panel">
                        <h2 class="text-xl font-bold text-emerald-400 mb-4">💵 Add Cash Expense & Files</h2>
                        <form method="POST" enctype="multipart/form-data" class="grid grid-cols-1 sm:grid-cols-3 gap-4">
                            <input type="hidden" name="form_type" value="add_cash_expense">
                            <div>
                                <label class="text-xs text-gray-400">Date</label>
                                <input type="date" name="date" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input">
                            </div>
                            <div>
                                <label class="text-xs text-gray-400">Pay To</label>
                                <input type="text" name="pay_to" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input" placeholder="Recipient Name">
                            </div>
                            <div>
                                <label class="text-xs text-gray-400">Description</label>
                                <input type="text" name="description" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input" placeholder="Expense purpose">
                            </div>
                            <div>
                                <label class="text-xs text-gray-400">Amount (₹)</label>
                                <input type="number" name="amount" value="0" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm dynamic-input">
                            </div>
                            <div class="sm:col-span-2">
                                <label class="text-xs text-gray-400">Upload Multiple Files (PDF/Images)</label>
                                <input type="file" name="attachments" multiple accept=".pdf,image/*" class="w-full mt-1 p-2 bg-gray-800 rounded-xl border border-gray-700 text-xs text-gray-300">
                            </div>
                            <div class="sm:col-span-3">
                                <button type="submit" class="w-full py-3 bg-emerald-500 hover:bg-emerald-600 font-bold text-gray-950 rounded-xl transition text-sm shadow-lg">Save Cash Expense</button>
                            </div>
                        </form>
                    </div>

                    <div class="bg-gray-900 rounded-2xl border border-gray-800 shadow-xl overflow-hidden card-panel">
                        <div class="p-4 border-b border-gray-800 flex flex-col sm:flex-row justify-between items-center gap-4">
                            <h3 class="font-bold dynamic-text">💵 Date-Wise Grouped Cash Expense Ledger</h3>
                            
                            <!-- Search & Export Controls for Cash Expense -->
                            <div class="flex flex-wrap items-center gap-2 w-full sm:w-auto justify-end">
                                <form method="GET" action="/" class="flex items-center gap-2 w-full sm:w-auto">
                                    <input type="hidden" name="action" value="cash_expenses">
                                    <input type="text" name="cash_search" value="{{ cash_search }}" placeholder="Search cash records..." class="px-3 py-1.5 bg-gray-800 border border-gray-700 rounded-xl text-xs text-white focus:outline-none focus:ring-2 focus:ring-emerald-400 w-full sm:w-64">
                                    <button type="submit" class="px-3 py-1.5 bg-emerald-500 text-gray-950 font-bold rounded-xl text-xs hover:bg-emerald-600 transition">Search</button>
                                    {% if cash_search %}
                                        <a href="/?action=cash_expenses" class="px-2 py-1.5 bg-gray-700 text-gray-300 rounded-xl text-xs hover:bg-gray-600">Clear</a>
                                    {% endif %}
                                </form>
                                <div class="flex items-center gap-1 border-l border-gray-700 pl-2">
                                    <a href="/export/cash_excel" class="px-3 py-1.5 bg-emerald-600 hover:bg-emerald-700 text-white font-bold rounded-xl text-xs transition flex items-center gap-1">📊 Excel</a>
                                    <a href="/export/cash_pdf" class="px-3 py-1.5 bg-red-600 hover:bg-red-700 text-white font-bold rounded-xl text-xs transition flex items-center gap-1">📄 PDF</a>
                                </div>
                            </div>
                        </div>

                        <div class="overflow-x-auto">
                            <table class="w-full text-left border-collapse text-xs">
                                <thead>
                                    <tr class="bg-gray-800/80 text-gray-300 uppercase tracking-wider font-mono">
                                        <th class="p-3 border-r border-gray-700">Sr #</th>
                                        <th class="p-3 border-r border-gray-700">Date</th>
                                        <th class="p-3 border-r border-gray-700">Pay To</th>
                                        <th class="p-3 border-r border-gray-700">Description</th>
                                        <th class="p-3 border-r border-gray-700">Amount</th>
                                        <th class="p-3 border-r border-gray-700">Files Attach</th>
                                        <th class="p-3 text-center">Actions</th>
                                    </tr>
                                </thead>
                                <tbody class="divide-y divide-gray-800 font-mono">
                                    {% set c_ns = namespace(current_date='', date_amt=0) %}
                                    {% for c in filtered_cash_expenses %}
                                        {% if c.date != c_ns.current_date %}
                                            {% if not loop.first %}
                                            <!-- Cash Date Subtotal Row -->
                                            <tr class="bg-emerald-950/40 font-bold border-t-2 border-emerald-500/40">
                                                <td colspan="4" class="p-2.5 text-right text-emerald-300">Subtotal for {{ c_ns.current_date }}:</td>
                                                <td colspan="3" class="p-2.5 text-emerald-400">₹{{ c_ns.date_amt }}</td>
                                            </tr>
                                            {% endif %}
                                            {% set c_ns.current_date = c.date %}
                                            {% set c_ns.date_amt = c.amount %}
                                            <!-- Cash Date Group Header -->
                                            <tr class="bg-emerald-900/40 text-emerald-400 font-bold">
                                                <td colspan="7" class="p-2.5 px-4 text-xs tracking-wider">📅 Date Group: {{ c.date or 'Unspecified Date' }}</td>
                                            </tr>
                                        {% else %}
                                            {% set c_ns.date_amt = c_ns.date_amt + c.amount %}
                                        {% endif %}

                                    <tr class="hover:bg-gray-800/35 transition">
                                        <td class="p-3 border-r border-gray-800">{{ loop.index }}</td>
                                        <td class="p-3 border-r border-gray-800 text-gray-300">{{ c.date }}</td>
                                        <td class="p-3 border-r border-gray-800 font-bold text-indigo-400">{{ c.pay_to }}</td>
                                        <td class="p-3 border-r border-gray-800 text-gray-200">{{ c.description }}</td>
                                        <td class="p-3 border-r border-gray-800 font-bold text-emerald-400">₹{{ c.amount }}</td>
                                        <td class="p-3 border-r border-gray-800">
                                            {% if c.attachments %}
                                                <div class="flex flex-col space-y-1">
                                                    {% for att in c.attachments %}
                                                    <button onclick="openFileViewer('data:application/octet-stream;base64,{{ att.file_data }}', '{{ att.file_name }}')" class="text-indigo-400 underline hover:text-indigo-300 text-[10px] text-left">📎 {{ att.file_name[:15] }}...</button>
                                                    {% endfor %}
                                                </div>
                                            {% else %}
                                                <span class="text-gray-500 text-[10px]">No Files</span>
                                            {% endif %}
                                        </td>
                                        <td class="p-3 text-center space-x-2 whitespace-nowrap">
                                            <button onclick='openEditCashModal({{ c|tojson|safe }})' class="text-blue-400 bg-blue-500/10 px-2 py-1 rounded text-[10px] font-bold cursor-pointer">Edit</button>
                                            <a href="/delete/cash_expense/{{ c.id }}" onclick="return confirm('Confirm delete cash expense?');" class="text-red-400 bg-red-500/10 px-2 py-1 rounded text-[10px] font-bold inline-block">Delete</a>
                                        </td>
                                    </tr>
                                        {% if loop.last and filtered_cash_expenses %}
                                        <!-- Last Cash Date Subtotal Row -->
                                        <tr class="bg-emerald-950/40 font-bold border-t-2 border-emerald-500/40">
                                            <td colspan="4" class="p-2.5 text-right text-emerald-300">Subtotal for {{ c_ns.current_date }}:</td>
                                            <td colspan="3" class="p-2.5 text-emerald-400">₹{{ c_ns.date_amt }}</td>
                                        </tr>
                                        {% endif %}
                                    {% endfor %}
                                </tbody>
                                <tfoot>
                                    <tr class="bg-gray-800/90 font-mono font-bold text-gray-200 border-t-2 border-gray-700">
                                        <td colspan="4" class="p-3 text-right uppercase tracking-wider text-emerald-400">Grand Total Sum:</td>
                                        <td colspan="3" class="p-3 text-emerald-400">₹{{ filtered_cash_total }}</td>
                                    </tr>
                                </tfoot>
                            </table>
                        </div>
                    </div>
                </div>
                {% endif %}

            </div>
        </div>
    </div>

    <!-- EDIT ATTENDANCE MODAL -->
    <div id="editModal" class="fixed inset-0 bg-black/70 hidden items-center justify-center p-4 z-50">
        <div class="bg-gray-900 border border-gray-800 rounded-2xl w-full max-w-xl p-6 shadow-2xl relative max-h-[90vh] overflow-y-auto">
            <button type="button" onclick="closeEditModal()" class="absolute top-4 right-4 text-gray-400 hover:text-white font-bold text-lg">✕</button>
            <h3 class="text-xl font-bold text-indigo-400 mb-4">Edit Attendance Record</h3>
            <form id="editForm" method="POST" enctype="multipart/form-data" class="space-y-4">
                <input type="hidden" name="form_type" value="edit_attendance">
                <input type="hidden" name="record_id" id="editRecordId">
                <div class="grid grid-cols-1 sm:grid-cols-2 gap-4">
                    <div><label class="text-xs text-gray-400">Date</label><input type="date" name="date" id="editDate" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm"></div>
                    <div><label class="text-xs text-gray-400">Name</label><input type="text" name="name" id="editName" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm"></div>
                    <div><label class="text-xs text-gray-400">Card Number</label><input type="text" name="card_number" id="editCard" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm"></div>
                    <div><label class="text-xs text-gray-400">Role</label><input type="text" name="role" id="editRole" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm"></div>
                    <div><label class="text-xs text-gray-400">In Time</label><input type="text" name="in_time" id="editIn" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm"></div>
                    <div><label class="text-xs text-gray-400">Out Time</label><input type="text" name="out_time" id="editOut" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm"></div>
                    <div><label class="text-xs text-gray-400">Signature</label><select name="signature" id="editSig" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm"><option value="Yes">Yes</option><option value="No">No</option><option value="Mistake">Mistake</option></select></div>
                    <div><label class="text-xs text-gray-400">Amount (₹)</label><input type="number" name="amount" id="editAmount" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm"></div>
                    <div><label class="text-xs text-gray-400">Conveyance (₹)</label><input type="number" name="conveyance" id="editConveyance" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm"></div>
                    <div class="sm:col-span-2"><label class="text-xs text-gray-400">Description</label><input type="text" name="description" id="editDesc" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm"></div>
                    
                    <div class="sm:col-span-2 bg-gray-800/50 p-3 rounded-xl border border-gray-700">
                        <label class="text-xs text-indigo-400 font-bold block mb-2">Manage Existing Uploaded Files (Uncheck to Remove):</label>
                        <div id="editAttachmentsList" class="space-y-2 max-h-36 overflow-y-auto"></div>
                    </div>

                    <div class="sm:col-span-2"><label class="text-xs text-gray-400">Upload More Files (PDF/Images)</label><input type="file" name="attachments" multiple accept=".pdf,image/*" class="w-full mt-1 p-2 bg-gray-800 rounded-xl border border-gray-700 text-xs text-gray-300"></div>
                </div>
                <button type="submit" class="w-full py-3 bg-indigo-500 hover:bg-indigo-600 text-gray-950 font-bold rounded-xl shadow-lg transition text-sm">Save Changes</button>
            </form>
        </div>
    </div>

    <!-- EDIT CASH EXPENSE MODAL -->
    <div id="editCashModal" class="fixed inset-0 bg-black/70 hidden items-center justify-center p-4 z-50">
        <div class="bg-gray-900 border border-gray-800 rounded-2xl w-full max-w-xl p-6 shadow-2xl relative max-h-[90vh] overflow-y-auto">
            <button type="button" onclick="closeEditCashModal()" class="absolute top-4 right-4 text-gray-400 hover:text-white font-bold text-lg">✕</button>
            <h3 class="text-xl font-bold text-emerald-400 mb-4">Edit Cash Expense Record</h3>
            <form id="editCashForm" method="POST" enctype="multipart/form-data" class="space-y-4">
                <input type="hidden" name="form_type" value="edit_cash_expense">
                <input type="hidden" name="record_id" id="editCashRecordId">
                <div class="grid grid-cols-1 sm:grid-cols-2 gap-4">
                    <div><label class="text-xs text-gray-400">Date</label><input type="date" name="date" id="editCashDate" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm"></div>
                    <div><label class="text-xs text-gray-400">Pay To</label><input type="text" name="pay_to" id="editCashPayTo" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm"></div>
                    <div class="sm:col-span-2"><label class="text-xs text-gray-400">Description</label><input type="text" name="description" id="editCashDesc" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm"></div>
                    <div class="sm:col-span-2"><label class="text-xs text-gray-400">Amount (₹)</label><input type="number" name="amount" id="editCashAmount" class="w-full mt-1 p-2.5 bg-gray-800 rounded-xl border border-gray-700 text-sm"></div>
                    
                    <div class="sm:col-span-2 bg-gray-800/50 p-3 rounded-xl border border-gray-700">
                        <label class="text-xs text-emerald-400 font-bold block mb-2">Manage Existing Uploaded Files (Uncheck to Remove):</label>
                        <div id="editCashAttachmentsList" class="space-y-2 max-h-36 overflow-y-auto"></div>
                    </div>

                    <div class="sm:col-span-2"><label class="text-xs text-gray-400">Upload More Files (PDF/Images)</label><input type="file" name="attachments" multiple accept=".pdf,image/*" class="w-full mt-1 p-2 bg-gray-800 rounded-xl border border-gray-700 text-xs text-gray-300"></div>
                </div>
                <button type="submit" class="w-full py-3 bg-emerald-500 hover:bg-emerald-600 text-gray-950 font-bold rounded-xl shadow-lg transition text-sm">Save Changes</button>
            </form>
        </div>
    </div>

    <!-- FILE ZOOM MODAL -->
    <div id="fileModal" class="fixed inset-0 bg-black/80 hidden items-center justify-center p-4 z-50">
        <div class="bg-gray-900 border border-gray-800 rounded-2xl w-full max-w-3xl p-4 shadow-2xl relative flex flex-col items-center">
            <button onclick="closeFileViewer()" class="absolute top-4 right-4 text-gray-400 hover:text-white font-bold text-lg z-10">✕</button>
            <h3 id="fileModalTitle" class="text-sm font-bold text-indigo-400 mb-2">Attachment Viewer</h3>
            <div id="fileViewerContainer" class="w-full max-h-[75vh] overflow-auto flex justify-center items-center p-2"></div>
        </div>
    </div>

    <script>
        document.addEventListener("DOMContentLoaded", () => {
            const savedTheme = localStorage.getItem('shivam_crt_theme');
            if (savedTheme === 'light') {
                applyLightTheme();
            }
        });

        function toggleTheme() {
            const body = document.getElementById('bodyTheme');
            if (body.classList.contains('bg-gray-950')) {
                applyLightTheme();
                localStorage.setItem('shivam_crt_theme', 'light');
            } else {
                applyDarkTheme();
                localStorage.setItem('shivam_crt_theme', 'dark');
            }
        }

        function applyLightTheme() {
            const body = document.getElementById('bodyTheme');
            const btn = document.getElementById('themeToggleBtn');
            const sidebar = document.getElementById('sidebarPanel');
            const header = document.getElementById('headerPanel');
            const cards = document.querySelectorAll('.card-panel');
            
            body.classList.remove('bg-gray-950', 'text-gray-100');
            body.classList.add('bg-gray-100', 'text-gray-900');
            if(btn) btn.innerText = "🌙 Dark Mode";
            if(sidebar) { sidebar.classList.remove('bg-gray-900', 'border-gray-800'); sidebar.classList.add('bg-white', 'border-gray-200'); }
            if(header) { header.classList.remove('bg-gray-900', 'border-gray-800'); header.classList.add('bg-white', 'border-gray-200'); }
            cards.forEach(c => { c.classList.remove('bg-gray-900', 'border-gray-800'); c.classList.add('bg-white', 'border-gray-200', 'shadow-md'); });
        }

        function applyDarkTheme() {
            const body = document.getElementById('bodyTheme');
            const btn = document.getElementById('themeToggleBtn');
            const sidebar = document.getElementById('sidebarPanel');
            const header = document.getElementById('headerPanel');
            const cards = document.querySelectorAll('.card-panel');
            
            body.classList.remove('bg-gray-100', 'text-gray-900');
            body.classList.add('bg-gray-950', 'text-gray-100');
            if(btn) btn.innerText = "☀️ Light Mode";
            if(sidebar) { sidebar.classList.add('bg-gray-900', 'border-gray-800'); sidebar.classList.remove('bg-white', 'border-gray-200'); }
            if(header) { header.classList.add('bg-gray-900', 'border-gray-800'); header.classList.remove('bg-white', 'border-gray-200'); }
            cards.forEach(c => { c.classList.add('bg-gray-900', 'border-gray-800'); c.classList.remove('bg-white', 'border-gray-200', 'shadow-md'); });
        }

        function openEditModal(item) {
            document.getElementById('editModal').classList.remove('hidden');
            document.getElementById('editModal').classList.add('flex');
            document.getElementById('editRecordId').value = item.id;
            document.getElementById('editDate').value = item.date || '';
            document.getElementById('editName').value = item.name || '';
            document.getElementById('editCard').value = item.card_number || '';
            document.getElementById('editRole').value = item.role || '';
            document.getElementById('editIn').value = item.in_time || '';
            document.getElementById('editOut').value = item.out_time || '';
            document.getElementById('editSig').value = item.signature || 'Yes';
            document.getElementById('editAmount').value = item.amount || 0;
            document.getElementById('editConveyance').value = item.conveyance || 0;
            document.getElementById('editDesc').value = item.description || '';

            const attListContainer = document.getElementById('editAttachmentsList');
            attListContainer.innerHTML = '';
            
            let files = item.attachments || [];
            if (files.length === 0) {
                attListContainer.innerHTML = '<p class="text-xs text-gray-400 italic">No files currently attached.</p>';
            } else {
                files.forEach((att, idx) => {
                    const div = document.createElement('div');
                    div.className = "flex items-center justify-between bg-gray-900 p-2 rounded-lg text-xs border border-gray-700";
                    div.innerHTML = `
                        <label class="flex items-center space-x-2 cursor-pointer truncate">
                            <input type="checkbox" name="keep_attachments" value="${idx}" checked class="rounded bg-gray-800 border-gray-700 text-indigo-500 focus:ring-indigo-400">
                            <span class="text-gray-300 truncate max-w-[200px]">📎 ${att.file_name}</span>
                        </label>
                        <button type="button" onclick="openFileViewer('data:application/octet-stream;base64,${att.file_data}', '${att.file_name}')" class="text-indigo-400 hover:underline px-2 py-1 bg-indigo-500/10 rounded">View</button>
                    `;
                    attListContainer.appendChild(div);
                });
            }
        }

        function closeEditModal() {
            document.getElementById('editModal').classList.remove('flex');
            document.getElementById('editModal').classList.add('hidden');
        }

        function openEditCashModal(item) {
            document.getElementById('editCashModal').classList.remove('hidden');
            document.getElementById('editCashModal').classList.add('flex');
            document.getElementById('editCashRecordId').value = item.id;
            document.getElementById('editCashDate').value = item.date || '';
            document.getElementById('editCashPayTo').value = item.pay_to || '';
            document.getElementById('editCashDesc').value = item.description || '';
            document.getElementById('editCashAmount').value = item.amount || 0;

            const attListContainer = document.getElementById('editCashAttachmentsList');
            attListContainer.innerHTML = '';
            
            let files = item.attachments || [];
            if (files.length === 0) {
                attListContainer.innerHTML = '<p class="text-xs text-gray-400 italic">No files currently attached.</p>';
            } else {
                files.forEach((att, idx) => {
                    const div = document.createElement('div');
                    div.className = "flex items-center justify-between bg-gray-900 p-2 rounded-lg text-xs border border-gray-700";
                    div.innerHTML = `
                        <label class="flex items-center space-x-2 cursor-pointer truncate">
                            <input type="checkbox" name="keep_attachments" value="${idx}" checked class="rounded bg-gray-800 border-gray-700 text-emerald-500 focus:ring-emerald-400">
                            <span class="text-gray-300 truncate max-w-[200px]">📎 ${att.file_name}</span>
                        </label>
                        <button type="button" onclick="openFileViewer('data:application/octet-stream;base64,${att.file_data}', '${att.file_name}')" class="text-emerald-400 hover:underline px-2 py-1 bg-emerald-500/10 rounded">View</button>
                    `;
                    attListContainer.appendChild(div);
                });
            }
        }

        function closeEditCashModal() {
            document.getElementById('editCashModal').classList.remove('flex');
            document.getElementById('editCashModal').classList.add('hidden');
        }

        function openFileViewer(dataUri, fileName) {
            const modal = document.getElementById('fileModal');
            const container = document.getElementById('fileViewerContainer');
            document.getElementById('fileModalTitle').innerText = "Viewing: " + fileName;
            container.innerHTML = '';
            
            if (fileName.toLowerCase().endsWith('.pdf')) {
                container.innerHTML = `<iframe src="${dataUri}" class="w-full h-[70vh] rounded border border-gray-700"></iframe>`;
            } else {
                container.innerHTML = `<img src="${dataUri}" class="max-w-full max-h-[70vh] object-contain rounded cursor-zoom-in" onclick="this.classList.toggle('scale-125')">`;
            }
            modal.classList.remove('hidden');
            modal.classList.add('flex');
        }

        function closeFileViewer() {
            document.getElementById('fileModal').classList.remove('flex');
            document.getElementById('fileModal').classList.add('hidden');
        }
    </script>
</body>
</html>
"""

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
